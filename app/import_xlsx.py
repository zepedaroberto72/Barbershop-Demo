"""Parse spreadsheet uploads into appointment rows (History → Import)."""
from __future__ import annotations

import csv
import re
from datetime import datetime
from io import BytesIO, StringIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook

# Canonical keys used after header normalization.
FIELD_ALIASES = {
    "id": ("id", "appointment id", "appt id"),
    "date": ("date", "day", "service date"),
    "time": ("time", "start time", "starts"),
    "starts_at": ("starts at", "start", "datetime", "date time", "starts_at"),
    "barber": ("barber", "tech", "technician", "stylist", "staff"),
    "customer": ("customer", "customer name", "name", "client"),
    "phone": ("phone", "mobile", "cell", "telephone"),
    "email": ("email", "e-mail"),
    "address": ("address", "location"),
    "service": ("service", "appliance", "service name"),
    "job_type": ("job type", "type", "service type"),
    "status": ("status",),
    "issue": ("issue", "request", "description"),
    "duration": ("duration (min)", "duration", "duration minutes", "minutes"),
    "travel": ("travel buffer (min)", "travel", "travel buffer", "travel minutes"),
    "notes": ("notes", "appointment notes", "job notes"),
    "customer_notes": ("customer notes", "client notes"),
}


def _norm_header(value: Any) -> str:
    raw = str(value or "").strip().lower()
    raw = raw.replace("_", " ").replace("-", " ")
    return re.sub(r"\s+", " ", raw)


def _header_map(headers: List[Any]) -> Dict[str, int]:
    """Map canonical field → column index."""
    by_norm = {_norm_header(h): i for i, h in enumerate(headers) if str(h or "").strip()}
    out: Dict[str, int] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in by_norm:
                out[field] = by_norm[alias]
                break
    return out


def _cell(row: List[Any], idx: Optional[int]) -> str:
    if idx is None or idx >= len(row):
        return ""
    val = row[idx]
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.isoformat(sep=" ", timespec="minutes")
    return str(val).strip()


def _parse_starts(date_s: str, time_s: str, combined: str) -> Optional[datetime]:
    combined = (combined or "").strip()
    if combined:
        for fmt in (
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %H:%M:%S",
            "%m/%d/%Y %H:%M",
            "%m/%d/%Y %I:%M %p",
            "%Y-%m-%dT%H:%M",
            "%Y-%m-%dT%H:%M:%S",
        ):
            try:
                return datetime.strptime(combined[:19], fmt)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(combined.replace("Z", ""))
        except ValueError:
            pass

    date_s = (date_s or "").strip()
    time_s = (time_s or "").strip() or "09:00"
    if not date_s:
        return None
    # openpyxl may give date-only datetime string
    if " " in date_s and not time_s:
        return _parse_starts("", "", date_s)

    d = None
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m-%d-%Y", "%d/%m/%Y"):
        try:
            d = datetime.strptime(date_s[:10], fmt).date()
            break
        except ValueError:
            continue
    if d is None:
        try:
            d = datetime.fromisoformat(date_s.replace("Z", "")).date()
        except ValueError:
            return None

    t = None
    for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M%p"):
        try:
            t = datetime.strptime(time_s.upper(), fmt).time()
            break
        except ValueError:
            continue
    if t is None:
        try:
            t = datetime.fromisoformat(f"2000-01-01T{time_s}").time()
        except ValueError:
            t = datetime.strptime("09:00", "%H:%M").time()
    return datetime.combine(d, t)


def _int_or(default: int, raw: str, *, minimum: int = 0) -> int:
    digits = re.sub(r"[^\d.]", "", raw or "")
    if not digits:
        return default
    try:
        return max(minimum, int(float(digits)))
    except ValueError:
        return default


def rows_from_workbook_bytes(content: bytes, filename: str = "") -> Tuple[List[str], List[Dict[str, Any]]]:
    """Return (errors, parsed_rows). Each row is a dict of canonical fields."""
    name = (filename or "").lower()
    errors: List[str] = []
    is_csv = name.endswith(".csv") or (
        not name.endswith((".xlsx", ".xlsm", ".xls"))
        and content[:1] not in (b"P", b"\x50")  # xlsx zip magic is PK
        and (b"," in content[:120] or b";" in content[:120])
    )
    if is_csv:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        reader = csv.reader(StringIO(text))
        table = [list(r) for r in reader]
    else:
        try:
            wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            return [f"Could not read spreadsheet: {exc}"], []
        ws = wb.active
        table = [list(row) for row in ws.iter_rows(values_only=True)]
        wb.close()

    # Drop fully empty trailing rows
    while table and all(c is None or str(c).strip() == "" for c in table[-1]):
        table.pop()
    if not table:
        return ["Spreadsheet is empty."], []

    headers = table[0]
    col = _header_map(headers)
    if "customer" not in col and "phone" not in col:
        return ["Missing required columns: Customer (or Phone)."], []
    if "date" not in col and "starts_at" not in col:
        return ["Missing required columns: Date (or Starts At)."], []

    parsed: List[Dict[str, Any]] = []
    for i, raw in enumerate(table[1:], start=2):
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        date_s = _cell(raw, col.get("date"))
        time_s = _cell(raw, col.get("time"))
        combined = _cell(raw, col.get("starts_at"))
        starts = _parse_starts(date_s, time_s, combined)
        if not starts:
            errors.append(f"Row {i}: invalid date/time.")
            continue
        customer = _cell(raw, col.get("customer")) or "Imported customer"
        phone = _cell(raw, col.get("phone")) or "000-0000"
        row = {
            "id": _cell(raw, col.get("id")),
            "starts_at": starts,
            "barber": _cell(raw, col.get("barber")),
            "customer": customer,
            "phone": phone,
            "email": _cell(raw, col.get("email")) or None,
            "address": _cell(raw, col.get("address")) or "Shop",
            "service": _cell(raw, col.get("service")) or "Haircut",
            "job_type": _cell(raw, col.get("job_type")) or _cell(raw, col.get("service")) or "Haircut",
            "status": (_cell(raw, col.get("status")) or "completed").strip().lower(),
            "issue": _cell(raw, col.get("issue")) or _cell(raw, col.get("service")) or "Imported",
            "duration": _int_or(30, _cell(raw, col.get("duration")), minimum=1),
            "travel": _int_or(0, _cell(raw, col.get("travel")), minimum=0),
            "notes": _cell(raw, col.get("notes")) or None,
            "customer_notes": _cell(raw, col.get("customer_notes")) or None,
            "row_num": i,
        }
        if row["status"] not in {"scheduled", "confirmed", "completed", "cancelled"}:
            row["status"] = "completed"
        parsed.append(row)
    return errors, parsed

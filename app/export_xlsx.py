"""Build Excel (.xlsx) workbooks for appointment and invoice exports."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models import Appointment

APPOINTMENT_HEADERS = [
    "ID",
    "Date",
    "Time",
    "Barber",
    "Customer",
    "Phone",
    "Email",
    "Address",
    "Service",
    "Job Type",
    "Status",
    "Issue",
    "Duration (min)",
    "Travel buffer (min)",
    "Notes",
    "Customer Notes",
    "Booked At",
]


def _style_header(ws, col_count: int) -> None:
    bold = Font(bold=True)
    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = bold
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def _autosize(ws, min_width: int = 10, max_width: int = 42) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        longest = 0
        for cell in col_cells:
            val = "" if cell.value is None else str(cell.value)
            longest = max(longest, len(val))
        ws.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


def _appointment_row(appt: Appointment) -> List[Any]:
    starts = appt.starts_at
    booked = getattr(appt, "created_at", None)
    return [
        appt.id,
        starts.strftime("%Y-%m-%d") if starts else "",
        starts.strftime("%H:%M") if starts else "",
        appt.technician.name if appt.technician else "",
        appt.customer.name if appt.customer else "",
        appt.customer.phone if appt.customer else "",
        (appt.customer.email if appt.customer else None) or "",
        appt.customer.address if appt.customer else "",
        appt.appliance or "",
        getattr(appt, "job_type", None) or "Haircut",
        appt.status or "",
        appt.issue or "",
        appt.duration_minutes,
        appt.travel_buffer_minutes,
        appt.notes or "",
        (appt.customer.notes if appt.customer else None) or "",
        booked.strftime("%Y-%m-%d %H:%M") if booked else "",
    ]


def _write_appointments_sheet(ws, appointments: Sequence[Appointment], title: str = "Appointments") -> None:
    ws.title = title[:31]
    ws.append(APPOINTMENT_HEADERS)
    _style_header(ws, len(APPOINTMENT_HEADERS))
    for appt in appointments:
        ws.append(_appointment_row(appt))
    _autosize(ws)


def _default_invoice_lines(appt: Appointment) -> List[Dict[str, Any]]:
    hours = max(1, round((appt.duration_minutes or 90) / 60 * 10) / 10)
    return [
        {
            "kind": "service",
            "description": f"{getattr(appt, 'job_type', None) or 'Repair'} — {appt.appliance}",
            "qty": hours,
            "rate": 125,
        },
        {"kind": "service", "description": "Service call / travel", "qty": 1, "rate": 49},
    ]


def _line_kind(line: Dict[str, Any]) -> str:
    kind = str(line.get("kind") or "service").strip().lower()
    return "material" if kind == "material" else "service"


def _append_charge_section(
    ws,
    *,
    title: str,
    headers: List[str],
    lines: List[Dict[str, Any]],
) -> float:
    ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(headers)
    header_row = ws.max_row
    for col in range(1, 5):
        ws.cell(row=header_row, column=col).font = Font(bold=True)
    subtotal = 0.0
    for line in lines:
        qty = float(line.get("qty") or 0)
        rate = float(line.get("rate") or 0)
        amount = qty * rate
        subtotal += amount
        # Columns: Description, Unit price/Price, Hours/Quantity, Total
        ws.append(
            [
                str(line.get("description") or ""),
                rate,
                qty,
                round(amount, 2),
            ]
        )
    return subtotal


def _write_current_invoice_sheet(
    ws,
    appt: Appointment,
    *,
    lines: Optional[List[Dict[str, Any]]] = None,
    tax_rate: float = 8.25,
    notes: Optional[str] = None,
    invoice_number: Optional[str] = None,
    invoice_date: Optional[str] = None,
) -> None:
    ws.title = "Current Invoice"
    starts = appt.starts_at
    number = invoice_number or (
        f"RD-{appt.id}-{starts.strftime('%Y%m%d')}" if starts else f"RD-{appt.id}"
    )
    date_str = invoice_date or datetime.now().strftime("%B %d, %Y")
    line_items = lines if lines is not None else _default_invoice_lines(appt)
    note_text = notes if notes is not None else (appt.notes or "Thank you for choosing Barber Shop.")

    meta = [
        ("Invoice #", number),
        ("Date", date_str),
        ("Customer", appt.customer.name if appt.customer else ""),
        ("Phone", appt.customer.phone if appt.customer else ""),
        ("Email", (appt.customer.email if appt.customer else None) or ""),
        ("Address", appt.customer.address if appt.customer else ""),
        ("Appliance", appt.appliance or ""),
        ("Job Type", getattr(appt, "job_type", None) or "Repair"),
        ("Status", appt.status or ""),
        ("Tech", appt.technician.name if appt.technician else ""),
        (
            "Service date",
            starts.strftime("%Y-%m-%d %H:%M") if starts else "",
        ),
        ("Notes", note_text),
    ]
    ws.append(["Field", "Value"])
    _style_header(ws, 2)
    for label, value in meta:
        ws.append([label, value])

    service_lines = [l for l in line_items if _line_kind(l) == "service"]
    material_lines = [l for l in line_items if _line_kind(l) == "material"]

    service_sub = _append_charge_section(
        ws,
        title="Service",
        headers=["Description", "Unit price", "Hours", "Total"],
        lines=service_lines,
    )
    materials_sub = _append_charge_section(
        ws,
        title="Materials",
        headers=["Description", "Price", "Quantity", "Total"],
        lines=material_lines,
    )

    subtotal = service_sub + materials_sub
    tax = subtotal * (float(tax_rate) / 100.0)
    total = subtotal + tax
    ws.append([])
    ws.append(["Service subtotal", "", "", round(service_sub, 2)])
    ws.append(["Materials subtotal", "", "", round(materials_sub, 2)])
    ws.append(["Subtotal", "", "", round(subtotal, 2)])
    ws.append([f"Tax ({tax_rate}%)", "", "", round(tax, 2)])
    ws.append(["Total", "", "", round(total, 2)])
    for row in range(ws.max_row - 4, ws.max_row + 1):
        ws.cell(row=row, column=1).font = Font(bold=True)
    _autosize(ws)


def build_past_workbook(appointments: Sequence[Appointment]) -> bytes:
    wb = Workbook()
    ws = wb.active
    _write_appointments_sheet(ws, appointments, title="History")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_invoice_workbook(
    jobs: Sequence[Appointment],
    *,
    current_job: Optional[Appointment] = None,
    lines: Optional[List[Dict[str, Any]]] = None,
    tax_rate: float = 8.25,
    notes: Optional[str] = None,
    invoice_number: Optional[str] = None,
    invoice_date: Optional[str] = None,
) -> bytes:
    wb = Workbook()
    ws_jobs = wb.active
    _write_appointments_sheet(ws_jobs, jobs, title="Invoice Jobs")
    if current_job is not None:
        ws_inv = wb.create_sheet()
        _write_current_invoice_sheet(
            ws_inv,
            current_job,
            lines=lines,
            tax_rate=tax_rate,
            notes=notes,
            invoice_number=invoice_number,
            invoice_date=invoice_date,
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
